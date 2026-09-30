import sqlite3
import smtplib
from email.mime.multipart import MIMEMultipart
from email.mime.base import MIMEBase
from email.mime.text import MIMEText
from email import encoders
import datetime
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

DB_PATH = '/opt/handover_app/handover.db'
SMTP_HOST = 'iss-365-mail1.itciss.com'
SMTP_PORT = 25
SENDER_EMAIL = 'shifthandover-noreply@itciss.com'
#RECIPIENTS = ['rkudipudi@itconvergence.com', 'klakkimsetti@itconvergence.com', 'rbeedilla@itconvergence.com']
RECIPIENTS = ['rkudipudi@itconvergence.com']

def generate_excel_report():
    today_str = datetime.datetime.now().strftime('%Y-%m-%d')
    excel_path = f'/tmp/Shift_Handover_Report_{today_str}.xlsx'

    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()

    # Fetch distinct dates in descending order (Latest date first)
    cursor.execute("""
        SELECT DISTINCT DATE(created_at) 
        FROM shift_handover 
        WHERE is_deleted = 0 
        ORDER BY DATE(created_at) DESC
    """)
    dates = [row[0] for row in cursor.fetchall() if row[0]]

    wb = openpyxl.Workbook()
    default_sheet = wb.active

    # Styling Definitions
    title_font = Font(name='Segoe UI', size=13, bold=True, color='FFFFFF')
    title_fill = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')

    header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='203764', end_color='203764', fill_type='solid')

    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    priority_fills = {
        'P1': PatternFill(start_color='F8D7DA', end_color='F8D7DA', fill_type='solid'),
        'P2': PatternFill(start_color='FFF3CD', end_color='FFF3CD', fill_type='solid'),
        'P3': PatternFill(start_color='E2E3E5', end_color='E2E3E5', fill_type='solid'),
        'P4': PatternFill(start_color='F8F9FA', end_color='F8F9FA', fill_type='solid')
    }

    status_fills = {
        'COMPLETED': PatternFill(start_color='D4EDDA', end_color='D4EDDA', fill_type='solid'),
        'IN_PROGRESS': PatternFill(start_color='CCE5FF', end_color='CCE5FF', fill_type='solid'),
        'PENDING': PatternFill(start_color='FFF3CD', end_color='FFF3CD', fill_type='solid')
    }

    headers = ["Priority", "Customer", "Ticket#", "Task Description", "Handover From", "Handover To", "Status", "Comments", "Timestamp"]

    if not dates:
        dates = [today_str]

    # Create Date-wise Separate Worksheets (Tabs)
    for idx, d_str in enumerate(dates):
        if idx == 0:
            ws = default_sheet
            ws.title = f"Date - {d_str}"
        else:
            ws = wb.create_sheet(title=f"Date - {d_str}")

        # Title Header Block
        ws.merge_cells('A1:I1')
        title_cell = ws['A1']
        title_cell.value = f"Hosting Team Daily Shift Handover Report ({d_str})"
        title_cell.font = title_font
        title_cell.fill = title_fill
        title_cell.alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions[1].height = 30

        # Table Column Headers
        ws.row_dimensions[2].height = 24
        for col_num, header in enumerate(headers, 1):
            c = ws.cell(row=2, column=col_num, value=header)
            c.font = header_font
            c.fill = header_fill
            c.alignment = Alignment(horizontal='center', vertical='center')
            c.border = thin_border

        # Fetch tasks for this specific date
        cursor.execute("""
            SELECT priority, shift_name, customer, ticket_no, task_description, resource_from, resource_to, status, comments, created_at 
            FROM shift_handover 
            WHERE is_deleted = 0 AND DATE(created_at) = ?
            ORDER BY shift_name ASC, id DESC
        """, (d_str,))
        tasks = cursor.fetchall()

        current_row = 3
        shifts = ["MORNING SHIFT TASKS", "AFTERNOON SHIFT TASKS", "NIGHT SHIFT TASKS"]
        
        for s in shifts:
            shift_tasks = [t for t in tasks if t[1] == s]
            if shift_tasks:
                ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=9)
                s_cell = ws.cell(row=current_row, column=1, value=s)
                s_cell.font = Font(name='Segoe UI', size=11, bold=True, color='1F4E78')
                s_cell.fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
                s_cell.alignment = Alignment(horizontal='center', vertical='center')
                ws.row_dimensions[current_row].height = 22
                current_row += 1

                for t in shift_tasks:
                    ws.row_dimensions[current_row].height = 20
                    # priority, customer, ticket_no, task_description, resource_from, resource_to, status, comments, created_at
                    row_data = [t[0], t[2], t[3], t[4], t[5], t[6], t[7], t[8], t[9]]
                    for col_num, val in enumerate(row_data, 1):
                        cell = ws.cell(row=current_row, column=col_num, value=val)
                        cell.font = Font(name='Segoe UI', size=10)
                        cell.border = thin_border
                        cell.alignment = Alignment(vertical='center')

                        if col_num in [1, 3, 5, 6, 7, 9]:
                            cell.alignment = Alignment(horizontal='center', vertical='center')

                        if col_num == 1 and val in priority_fills:
                            cell.fill = priority_fills[val]
                            cell.font = Font(name='Segoe UI', size=10, bold=True)
                        elif col_num == 7 and val in status_fills:
                            cell.fill = status_fills[val]
                            cell.font = Font(name='Segoe UI', size=10, bold=True)

                    current_row += 1

        # Auto-adjust Column Widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.row == 1:
                    continue
                if cell.value:
                    val_str = str(cell.value)
                    if len(val_str) > max_len:
                        max_len = len(val_str)
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 45)

    conn.close()
    wb.save(excel_path)
    return excel_path

def send_email():
    try:
        excel_path = generate_excel_report()
        today_formatted = datetime.datetime.now().strftime('%d-%b-%Y')
        
        msg = MIMEMultipart()
        msg['From'] = SENDER_EMAIL
        msg['To'] = ", ".join(RECIPIENTS)
        msg['Subject'] = f"Hosting Team Daily Shift Handover Report - {today_formatted}"

        body = f"""
Dear Team,

Please find attached the Daily Shift Handover Report for {today_formatted}.

This workbook includes Date-wise Separate Worksheets (Tabs) at the bottom for easy tracking across different shifts and dates.

Best Regards,
Hosting Operations Team
        """
        msg.attach(MIMEText(body, 'plain'))

        with open(excel_path, 'rb') as f:
            part = MIMEBase('application', 'octet-stream')
            part.set_payload(f.read())
            encoders.encode_base64(part)
            part.add_header('Content-Disposition', f'attachment; filename="{os.path.basename(excel_path)}"')
            msg.attach(part)

        server = smtplib.SMTP(SMTP_HOST, SMTP_PORT, timeout=15)
        server.sendmail(SENDER_EMAIL, RECIPIENTS, msg.as_string())
        server.quit()
        print("Daily report with Date-wise Tabs sent successfully!")
    except Exception as e:
        print(f"Error sending daily report email: {e}")

if __name__ == '__main__':
    send_email()
