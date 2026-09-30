import os
import sqlite3
import threading
import secrets
import string
import pandas as pd

from flask import Flask, render_template, request, redirect, url_for, session, flash, send_file
from flask_mail import Mail, Message
from datetime import datetime
from werkzeug.security import generate_password_hash, check_password_hash

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

app = Flask(__name__)
app.secret_key = "hosting_team_shift_handover_secret_key"
DB_PATH = "/opt/handover_app/handover.db"
PORTAL_URL = "http://cms-365-stigtst09.itciss.com"

app.config['MAIL_SERVER'] = 'iss-365-mail1.itciss.com'
app.config['MAIL_PORT'] = 25
app.config['MAIL_USE_TLS'] = False
app.config['MAIL_USE_SSL'] = False
app.config['MAIL_USERNAME'] = None
app.config['MAIL_PASSWORD'] = None
app.config['MAIL_DEFAULT_SENDER'] = ('ITConvergence Handover Portal', 'etkupro@cms-365-olam01.itciss.com')

mail = Mail(app)

def generate_random_temp_password():
    chars = string.ascii_letters + string.digits
    random_str = ''.join(secrets.choice(chars) for _ in range(5))
    return f"ITC@{random_str}!"

def get_db_connection():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn

def init_db_columns():
    try:
        conn = get_db_connection()
        conn.execute("ALTER TABLE tasks ADD COLUMN handover_notes TEXT DEFAULT '';")
        conn.commit()
    except Exception:
        pass

    try:
        conn = get_db_connection()
        conn.execute("ALTER TABLE tasks ADD COLUMN original_created_at TIMESTAMP;")
        conn.commit()
        conn.execute("UPDATE tasks SET original_created_at = created_at WHERE original_created_at IS NULL;")
        conn.commit()
        conn.close()
    except Exception:
        pass

init_db_columns()

def auto_carry_forward_pending_tasks():
    try:
        today_str = datetime.now().strftime('%Y-%m-%d')
        conn = get_db_connection()
        cursor = conn.cursor()

        past_pending_tasks = cursor.execute('''
            SELECT * FROM tasks 
            WHERE status NOT IN ('COMPLETED', 'CARRIED_FORWARD') 
              AND DATE(created_at) < DATE(?)
        ''', (today_str,)).fetchall()
        
        for task in past_pending_tasks:
            ticket_no = task['ticket_no']
            team_name = task['team_name']
            
            already_copied_today = cursor.execute('''
                SELECT id FROM tasks 
                WHERE ticket_no = ? AND DATE(created_at) = DATE(?) AND team_name = ?
            ''', (ticket_no, today_str, team_name)).fetchone()
            
            if not already_copied_today:
                task_keys = task.keys()
                orig_created = task['original_created_at'] if ('original_created_at' in task_keys and task['original_created_at']) else task['created_at']

                cursor.execute('''
                    INSERT INTO tasks (shift_name, customer, ticket_no, task_description, resource_from, resource_to, status, comments, priority, team_name, handover_notes, original_created_at)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ''', (
                    task['shift_name'], 
                    task['customer'], 
                    ticket_no, 
                    task['task_description'], 
                    task['resource_from'], 
                    task['resource_to'], 
                    task['status'], 
                    task['comments'], 
                    task['priority'], 
                    team_name, 
                    task['handover_notes'],
                    orig_created
                ))
                
                cursor.execute("UPDATE tasks SET status = 'CARRIED_FORWARD' WHERE id = ?", (task['id'],))
                
        conn.commit()
        conn.close()
    except Exception as e:
        print(f"Auto Carry Forward Error: {e}")

def send_async_email(app_obj, msg):
    with app_obj.app_context():
        try:
            mail.send(msg)
            print("Background Email Sent Successfully!")
        except Exception as e:
            print(f"Background Email Error: {e}")

def send_p1p2_alert_email(ticket_no, customer, priority, shift_name, task_desc, resource_from, resource_to, assignee_email, team_name):
    target_email = assignee_email if assignee_email else "rkudipudi@itconvergence.com"
    subject = f"🚨 [{priority} CRITICAL ALERT] New Task #{ticket_no} - {customer} ({team_name} Team)"
    body = f"""
Hello {resource_to},

A Critical {priority} Handover Task has been assigned to you.

📌 Task Details:
-------------------------------------------
- Team: {team_name}
- Priority: {priority}
- Customer: {customer}
- Ticket / SCTASK: {ticket_no}
- Shift: {shift_name}
- Handed Over By: {resource_from}

🌐 Handover Portal Link:
{PORTAL_URL}

📝 Instructions / Description:
{task_desc}

Please review the Handover Dashboard immediately.
    """
    msg = Message(subject, recipients=[target_email], body=body)
    threading.Thread(target=send_async_email, args=(app, msg)).start()

def send_user_welcome_email(email, username, full_name, temp_password, team_name):
    target_email = email if email and email != 'N/A' else "rkudipudi@itconvergence.com"
    subject = f"🔐 Access Credentials for ITConvergence Shift Handover Portal ({team_name} Team)"
    body = f"""
Hello {full_name},

An account credential update/creation has occurred for you in the Shift Handover Portal.

📌 Login Details:
-------------------------------------------
- Portal Team: {team_name} Team
- Username: {username}
- Temporary Password: {temp_password}

🌐 Portal Login Link:
{PORTAL_URL}/login

Note: You will be required to change your password upon logging in with this temporary password.
    """
    msg = Message(subject, recipients=[target_email], body=body)
    threading.Thread(target=send_async_email, args=(app, msg)).start()

def log_audit_action(task_id, action_type, details, username, team_name):
    try:
        conn = get_db_connection()
        conn.execute(
            "INSERT INTO audit_logs (task_id, action_type, details, performed_by) VALUES (?, ?, ?, ?)",
            (task_id, action_type, f"[{team_name}] {details}", username)
        )
        conn.commit()
        conn.close()
    except Exception as e:
        print(f"Audit Log Error: {e}")

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        username = request.form['username'].strip()
        password = request.form['password'].strip()

        try:
            conn = get_db_connection()
            user = conn.execute("SELECT * FROM users WHERE username = ?", (username,)).fetchone()
            conn.close()

            if user:
                user_keys = user.keys()
                stored_pw = user['password'] if 'password' in user_keys else user['password_hash']

                pw_matched = False
                try:
                    pw_matched = check_password_hash(stored_pw, password)
                except Exception:
                    pw_matched = (stored_pw == password)

                if pw_matched:
                    session['user_id'] = user['id']
                    session['username'] = user['username']
                    session['role'] = user['role']
                    session['team_name'] = user['team_name'] if ('team_name' in user_keys and user['team_name']) else 'Hosting'
                    session['selected_team'] = session['team_name']

                    force_change = user['force_password_change'] if ('force_password_change' in user_keys and user['force_password_change'] is not None) else 0
                    if force_change == 1:
                        session['must_change_password'] = True
                        return redirect(url_for('change_password'))

                    auto_carry_forward_pending_tasks()
                    return redirect(url_for('index'))

            flash("Invalid Username or Password!", "danger")
        except Exception as e:
            flash(f"Login Error: {str(e)}", "danger")

    return render_template('login.html')

@app.route('/change_password', methods=['GET', 'POST'])
def change_password():
    if 'user_id' not in session:
        return redirect(url_for('login'))

    if request.method == 'POST':
        new_password = request.form['new_password'].strip()
        confirm_password = request.form['confirm_password'].strip()

        if new_password != confirm_password:
            flash("Passwords do not match!", "danger")
            return render_template('change_password.html')

        hashed_pw = generate_password_hash(new_password)
        conn = get_db_connection()
        conn.execute("UPDATE users SET password = ?, force_password_change = 0 WHERE id = ?", (hashed_pw, session['user_id']))
        conn.commit()
        conn.close()

        session.pop('must_change_password', None)
        flash("Password updated successfully! Welcome to the portal.", "success")
        return redirect(url_for('index'))

    return render_template('change_password.html')

@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('login'))

@app.route('/switch_team', methods=['POST'])
def switch_team():
    if 'user_id' not in session or session.get('role') not in ['ADMIN', 'LEAD']:
        return redirect(url_for('index'))

    selected_team = request.form.get('selected_team')
    if selected_team:
        session['selected_team'] = selected_team
    return redirect(url_for('index'))

@app.route('/')
def index():
    if 'user_id' not in session:
        return redirect(url_for('login'))

    if session.get('must_change_password'):
        return redirect(url_for('change_password'))

    auto_carry_forward_pending_tasks()

    conn = get_db_connection()

    try:
        teams_list = [row['team_name'] for row in conn.execute("SELECT team_name FROM teams ORDER BY team_name ASC").fetchall()]
    except Exception:
        teams_list = ['Hosting', 'DBA', 'Linux/Unix', 'Windows/Cloud', 'Network', 'DevSecOps']

    if session.get('role') in ['ADMIN', 'LEAD']:
        active_team = session.get('selected_team', session.get('team_name', 'Hosting'))
    else:
        active_team = session.get('team_name', 'Hosting')

    current_date = datetime.now().strftime('%Y-%m-%d')

    try:
        if active_team == 'ALL_TEAMS' and session.get('role') in ['ADMIN', 'LEAD']:
            tasks = conn.execute("SELECT * FROM tasks ORDER BY id DESC").fetchall()
            total_count = conn.execute("SELECT COUNT(*) FROM tasks").fetchone()[0]
            pending_count = conn.execute("SELECT COUNT(*) FROM tasks WHERE status NOT IN ('COMPLETED', 'CARRIED_FORWARD')").fetchone()[0]
            completed_count = conn.execute("SELECT COUNT(*) FROM tasks WHERE status = 'COMPLETED'").fetchone()[0]
            morning_count = conn.execute("SELECT COUNT(*) FROM tasks WHERE shift_name = 'MORNING SHIFT TASKS'").fetchone()[0]
            afternoon_count = conn.execute("SELECT COUNT(*) FROM tasks WHERE shift_name = 'AFTERNOON SHIFT TASKS'").fetchone()[0]
            night_count = conn.execute("SELECT COUNT(*) FROM tasks WHERE shift_name = 'NIGHT SHIFT TASKS'").fetchone()[0]
            critical_p12_count = conn.execute("SELECT COUNT(*) FROM tasks WHERE priority IN ('P1', 'P2') AND status NOT IN ('COMPLETED', 'CARRIED_FORWARD')").fetchone()[0]
        else:
            tasks = conn.execute("SELECT * FROM tasks WHERE team_name = ? ORDER BY id DESC", (active_team,)).fetchall()
            total_count = conn.execute("SELECT COUNT(*) FROM tasks WHERE team_name = ?", (active_team,)).fetchone()[0]
            pending_count = conn.execute("SELECT COUNT(*) FROM tasks WHERE team_name = ? AND status NOT IN ('COMPLETED', 'CARRIED_FORWARD')", (active_team,)).fetchone()[0]
            completed_count = conn.execute("SELECT COUNT(*) FROM tasks WHERE team_name = ? AND status = 'COMPLETED'", (active_team,)).fetchone()[0]
            morning_count = conn.execute("SELECT COUNT(*) FROM tasks WHERE team_name = ? AND shift_name = 'MORNING SHIFT TASKS'", (active_team,)).fetchone()[0]
            afternoon_count = conn.execute("SELECT COUNT(*) FROM tasks WHERE team_name = ? AND shift_name = 'AFTERNOON SHIFT TASKS'", (active_team,)).fetchone()[0]
            night_count = conn.execute("SELECT COUNT(*) FROM tasks WHERE team_name = ? AND shift_name = 'NIGHT SHIFT TASKS'", (active_team,)).fetchone()[0]
            critical_p12_count = conn.execute("SELECT COUNT(*) FROM tasks WHERE team_name = ? AND priority IN ('P1', 'P2') AND status NOT IN ('COMPLETED', 'CARRIED_FORWARD')", (active_team,)).fetchone()[0]
    except Exception as e:
        print(f"Tasks Query Error: {e}")
        tasks = []
        total_count = pending_count = completed_count = morning_count = afternoon_count = night_count = critical_p12_count = 0

    try:
        if active_team == 'ALL_TEAMS' and session.get('role') in ['ADMIN', 'LEAD']:
            audit_logs = conn.execute("SELECT * FROM audit_logs ORDER BY id DESC LIMIT 100").fetchall()
        else:
            team_pattern = f"[{active_team}]%"
            audit_logs = conn.execute("SELECT * FROM audit_logs WHERE details LIKE ? ORDER BY id DESC LIMIT 100", (team_pattern,)).fetchall()
    except Exception as e:
        print(f"Audit Logs Query Error: {e}")
        audit_logs = []

    try:
        users_list = conn.execute("SELECT * FROM users ORDER BY id ASC").fetchall()
    except Exception:
        users_list = []

    conn.close()

    return render_template(
        'index.html',
        tasks=tasks,
        total_count=total_count,
        pending_count=pending_count,
        completed_count=completed_count,
        morning_count=morning_count,
        afternoon_count=afternoon_count,
        night_count=night_count,
        critical_p12_count=critical_p12_count,
        current_date=current_date,
        current_user=session.get('username', 'User'),
        user_role=session.get('role', 'ENGINEER'),
        user_team=session.get('team_name', 'Hosting'),
        active_team=active_team,
        teams_list=teams_list,
        audit_logs=audit_logs,
        users_list=users_list
    )

@app.route('/add_shift_task', methods=['POST'])
def add_shift_task():
    if 'user_id' not in session:
        return redirect(url_for('login'))

    try:
        shift_name = request.form.get('shift_name', 'MORNING SHIFT TASKS')
        customer = request.form.get('customer', 'N/A')
        ticket_no = request.form.get('ticket_no', 'N/A')
        task_description = request.form.get('task_description', '')
        resource_from = session.get('username', 'User')
        resource_to = request.form.get('resource_to', 'Unassigned')
        assignee_email = request.form.get('assignee_email', 'rkudipudi@itconvergence.com')
        status = request.form.get('status', 'PENDING')
        comments = request.form.get('comments', '')
        priority = request.form.get('priority', 'P3')
        handover_notes = request.form.get('handover_notes', '')
        now_ts = datetime.now().strftime('%Y-%m-%d %H:%M:%S')

        team_name = session.get('selected_team', session.get('team_name', 'Hosting'))
        if team_name == 'ALL_TEAMS':
            team_name = session.get('team_name', 'Hosting')

        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute('''
            INSERT INTO tasks (shift_name, customer, ticket_no, task_description, resource_from, resource_to, status, comments, priority, team_name, handover_notes, original_created_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ''', (shift_name, customer, ticket_no, task_description, resource_from, resource_to, status, comments, priority, team_name, handover_notes, now_ts))

        new_task_id = cursor.lastrowid
        conn.commit()
        conn.close()

        if priority in ['P1', 'P2']:
            send_p1p2_alert_email(ticket_no, customer, priority, shift_name, task_description, resource_from, resource_to, assignee_email, team_name)

        log_audit_action(new_task_id, "TASK_CREATED", f"Ticket #{ticket_no} created ({priority})", session['username'], team_name)
        flash("Shift handover task submitted successfully!", "success")
    except Exception as e:
        print(f"Task Creation Error: {e}")
        flash(f"Error creating task: {str(e)}", "danger")

    return redirect(url_for('index'))

@app.route('/edit_task', methods=['POST'])
def edit_task():
    if 'user_id' not in session:
        return redirect(url_for('login'))

    task_id = request.form['task_id']
    shift_name = request.form['shift_name']
    customer = request.form['customer']
    ticket_no = request.form['ticket_no']
    task_description = request.form['task_description']
    resource_from = session['username']
    resource_to = request.form['resource_to']
    assignee_email = request.form.get('assignee_email', 'rkudipudi@itconvergence.com')
    new_status = request.form['status']
    comments = request.form.get('comments', '')
    priority = request.form.get('priority', 'P3')
    handover_notes = request.form.get('handover_notes', '')

    conn = get_db_connection()
    
    # Protect status if task is already CARRIED_FORWARD
    current_task = conn.execute("SELECT status FROM tasks WHERE id = ?", (task_id,)).fetchone()
    if current_task and current_task['status'] == 'CARRIED_FORWARD':
        new_status = 'CARRIED_FORWARD'

    conn.execute('''
        UPDATE tasks
        SET shift_name=?, customer=?, ticket_no=?, task_description=?, resource_from=?, resource_to=?, status=?, comments=?, priority=?, handover_notes=?
        WHERE id=?
    ''', (shift_name, customer, ticket_no, task_description, resource_from, resource_to, new_status, comments, priority, handover_notes, task_id))
    conn.commit()
    conn.close()

    if priority in ['P1', 'P2']:
        send_p1p2_alert_email(ticket_no, customer, priority, shift_name, task_description, resource_from, resource_to, assignee_email, session.get('team_name', 'Hosting'))

    log_audit_action(task_id, "TASK_UPDATED", f"Updated details for Ticket #{ticket_no}", session['username'], session.get('team_name', 'Hosting'))
    flash("Task updated successfully!", "success")
    return redirect(url_for('index'))

@app.route('/update_notes', methods=['POST'])
def update_notes():
    if 'user_id' not in session:
        return redirect(url_for('login'))

    task_id = request.form.get('task_id')
    handover_notes = request.form.get('handover_notes', '').strip()

    conn = get_db_connection()
    conn.execute("UPDATE tasks SET handover_notes = ? WHERE id = ?", (handover_notes, task_id))
    conn.commit()
    conn.close()

    log_audit_action(task_id, "NOTES_UPDATED", f"Updated handover notes for Task #{task_id}", session['username'], session.get('team_name', 'Hosting'))
    flash("Handover notes updated successfully!", "success")
    return redirect(url_for('index'))

@app.route('/update_status/<int:task_id>/<string:new_status>')
def update_status(task_id, new_status):
    if 'user_id' not in session:
        return redirect(url_for('login'))

    conn = get_db_connection()
    conn.execute("UPDATE tasks SET status = ? WHERE id = ?", (new_status, task_id))
    conn.commit()
    conn.close()

    log_audit_action(task_id, "STATUS_CHANGED", f"Marked task #{task_id} as {new_status}", session['username'], session.get('team_name', 'Hosting'))
    return redirect(url_for('index'))

@app.route('/delete_task/<int:task_id>')
def delete_task(task_id):
    if 'user_id' not in session:
        return redirect(url_for('login'))

    conn = get_db_connection()
    conn.execute("DELETE FROM tasks WHERE id = ?", (task_id,))
    conn.commit()
    conn.close()

    log_audit_action(task_id, "TASK_DELETED", f"Deleted task ID #{task_id}", session['username'], session.get('team_name', 'Hosting'))
    return redirect(url_for('index'))

@app.route('/add_team', methods=['POST'])
def add_team():
    if 'user_id' not in session or session.get('role') != 'ADMIN':
        return redirect(url_for('index'))

    new_team_name = request.form['new_team_name'].strip()
    if new_team_name:
        conn = get_db_connection()
        try:
            conn.execute("INSERT INTO teams (team_name) VALUES (?)", (new_team_name,))
            conn.commit()
            flash(f"New Team '{new_team_name}' created successfully!", "success")
        except sqlite3.IntegrityError:
            flash(f"Team '{new_team_name}' already exists!", "danger")
        conn.close()

    return redirect(url_for('index') + '#teams-tab')

@app.route('/delete_team/<int:team_id>')
def delete_team(team_id):
    if 'user_id' not in session or session.get('role') != 'ADMIN':
        return redirect(url_for('index'))

    conn = get_db_connection()
    conn.execute("DELETE FROM teams WHERE id = ?", (team_id,))
    conn.commit()
    conn.close()

    return redirect(url_for('index') + '#teams-tab')

@app.route('/add_user', methods=['POST'])
def add_user():
    if 'user_id' not in session or session.get('role') != 'ADMIN':
        return redirect(url_for('index'))

    username = request.form['username'].strip()
    full_name = request.form['full_name'].strip()
    email = request.form['email'].strip()
    role = request.form['role']
    team_name = request.form['team_name']

    temp_password = generate_random_temp_password()
    hashed_pw = generate_password_hash(temp_password)

    conn = get_db_connection()
    try:
        conn.execute(
            "INSERT INTO users (username, full_name, email, role, password, team_name, force_password_change) VALUES (?, ?, ?, ?, ?, ?, 1)",
            (username, full_name, email, role, hashed_pw, team_name)
        )
        conn.commit()

        send_user_welcome_email(email, username, full_name, temp_password, team_name)

        flash(f"User '{username}' created for Team '{team_name}'! Temp Password: {temp_password}", "success")
    except sqlite3.IntegrityError:
        flash("Username already exists!", "danger")
    conn.close()

    return redirect(url_for('index') + '#users-tab')

@app.route('/edit_user', methods=['POST'])
def edit_user():
    if 'user_id' not in session or session.get('role') != 'ADMIN':
        return redirect(url_for('index'))

    user_id = request.form['user_id']
    full_name = request.form['full_name'].strip()
    email = request.form['email'].strip()
    role = request.form['role']
    team_name = request.form['team_name']

    conn = get_db_connection()
    conn.execute("UPDATE users SET full_name = ?, email = ?, role = ?, team_name = ? WHERE id = ?", (full_name, email, role, team_name, user_id))
    conn.commit()
    conn.close()

    flash("User updated!", "success")
    return redirect(url_for('index') + '#users-tab')

@app.route('/reset_user_password', methods=['POST'])
def reset_user_password():
    if 'user_id' not in session or session.get('role') != 'ADMIN':
        return redirect(url_for('index'))

    user_id = request.form['user_id']

    temp_password = generate_random_temp_password()
    hashed_pw = generate_password_hash(temp_password)

    conn = get_db_connection()
    user = conn.execute("SELECT * FROM users WHERE id = ?", (user_id,)).fetchone()
    if user:
        conn.execute("UPDATE users SET password = ?, force_password_change = 1 WHERE id = ?", (hashed_pw, user_id))
        conn.commit()

        user_keys = user.keys()
        email = user['email'] if ('email' in user_keys and user['email']) else "rkudipudi@itconvergence.com"
        team_name = user['team_name'] if ('team_name' in user_keys and user['team_name']) else 'Hosting'

        send_user_welcome_email(email, user['username'], user['full_name'], temp_password, team_name)
        flash(f"Password reset for '{user['username']}'! Temp Password: {temp_password}", "success")
    conn.close()

    return redirect(url_for('index') + '#users-tab')

@app.route('/delete_user/<int:user_id>')
def delete_user(user_id):
    if 'user_id' not in session or session.get('role') != 'ADMIN':
        return redirect(url_for('index'))

    conn = get_db_connection()
    conn.execute("DELETE FROM users WHERE id = ?", (user_id,))
    conn.commit()
    conn.close()

    flash("User account removed!", "info")
    return redirect(url_for('index') + '#users-tab')

@app.route('/export_shift_excel')
def export_shift_excel():
    if 'user_id' not in session:
        return redirect(url_for('login'))

    active_team = session.get('selected_team', session.get('team_name', 'Hosting'))
    conn = get_db_connection()

    if active_team == 'ALL_TEAMS' and session.get('role') in ['ADMIN', 'LEAD']:
        df = pd.read_sql_query("SELECT *, DATE(created_at) as log_date FROM tasks ORDER BY id DESC", conn)
    else:
        df = pd.read_sql_query("SELECT *, DATE(created_at) as log_date FROM tasks WHERE team_name = ? ORDER BY id DESC", conn, params=(active_team,))

    conn.close()

    wb = Workbook()
    wb.remove(wb.active)

    font_banner = Font(name='Calibri', size=14, bold=True, color='FFFFFF')
    fill_banner = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')

    font_col_hdr = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    fill_col_hdr = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')

    font_shift_hdr = Font(name='Calibri', size=11, bold=True, color='1F4E78')
    fill_shift_hdr = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')

    font_p1 = Font(name='Calibri', size=11, bold=True, color='9C0006')
    fill_p1 = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')

    font_p2 = Font(name='Calibri', size=11, bold=True, color='9C6500')
    fill_p2 = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')

    fill_pending = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')
    fill_completed = PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid')

    align_center = Alignment(horizontal='center', vertical='center')
    align_left = Alignment(horizontal='left', vertical='center')
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    headers = ['Priority', 'Customer', 'Ticket#', 'Task Description', 'Handover From', 'Handover To', 'Status', 'Handover Notes', 'Comments', 'Timestamp']
    shifts_order = ['MORNING SHIFT TASKS', 'AFTERNOON SHIFT TASKS', 'NIGHT SHIFT TASKS']

    if df.empty:
        ws = wb.create_sheet(title="Date - No Data")
        ws.merge_cells('A1:J1')
        ws['A1'] = f"{active_team} Team Daily Shift Handover Report"
        ws['A1'].font = font_banner
        ws['A1'].fill = fill_banner
        ws['A1'].alignment = align_center
    else:
        grouped = df.groupby('log_date')
        for log_date, group in grouped:
            sheet_title = f"Date - {log_date}"
            ws = wb.create_sheet(title=sheet_title[:31])

            ws.merge_cells('A1:J1')
            ws['A1'] = f"{active_team} Team Daily Shift Handover Report ({log_date})"
            ws['A1'].font = font_banner
            ws['A1'].fill = fill_banner
            ws['A1'].alignment = align_center
            ws.row_dimensions[1].height = 28

            for col_idx, h in enumerate(headers, 1):
                cell = ws.cell(row=2, column=col_idx, value=h)
                cell.font = font_col_hdr
                cell.fill = fill_col_hdr
                cell.alignment = align_center
                cell.border = thin_border
            ws.row_dimensions[2].height = 24

            current_row = 3

            for shift in shifts_order:
                shift_tasks = group[group['shift_name'] == shift]
                if not shift_tasks.empty:
                    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=10)
                    s_cell = ws.cell(row=current_row, column=1, value=shift)
                    s_cell.font = font_shift_hdr
                    s_cell.fill = fill_shift_hdr
                    s_cell.alignment = align_center
                    ws.row_dimensions[current_row].height = 20
                    current_row += 1

                    for _, row in shift_tasks.iterrows():
                        ws.cell(row=current_row, column=1, value=row.get('priority', 'P3')).alignment = align_center
                        ws.cell(row=current_row, column=2, value=row.get('customer', 'N/A')).alignment = align_left
                        ws.cell(row=current_row, column=3, value=row.get('ticket_no', 'N/A')).alignment = align_center
                        ws.cell(row=current_row, column=4, value=row.get('task_description', '')).alignment = align_left
                        ws.cell(row=current_row, column=5, value=row.get('resource_from', '')).alignment = align_center
                        ws.cell(row=current_row, column=6, value=row.get('resource_to', '')).alignment = align_center
                        ws.cell(row=current_row, column=7, value=row.get('status', 'PENDING')).alignment = align_center
                        ws.cell(row=current_row, column=8, value=row.get('handover_notes', '')).alignment = align_left
                        ws.cell(row=current_row, column=9, value=row.get('comments', '')).alignment = align_left
                        ws.cell(row=current_row, column=10, value=str(row.get('created_at', ''))).alignment = align_center

                        for c in range(1, 11):
                            cell = ws.cell(row=current_row, column=c)
                            cell.border = thin_border

                        p_cell = ws.cell(row=current_row, column=1)
                        if p_cell.value == 'P1':
                            p_cell.font = font_p1
                            p_cell.fill = fill_p1
                        elif p_cell.value == 'P2':
                            p_cell.font = font_p2
                            p_cell.fill = fill_p2

                        st_cell = ws.cell(row=current_row, column=7)
                        if st_cell.value == 'COMPLETED':
                            st_cell.fill = fill_completed
                            st_cell.font = Font(bold=True, color='276A3C')
                        else:
                            st_cell.fill = fill_pending
                            st_cell.font = Font(bold=True, color='B25900')

                        ws.row_dimensions[current_row].height = 22
                        current_row += 1

            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    filename = f"/tmp/Shift_Handover_Report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    wb.save(filename)
    return send_file(filename, as_attachment=True)

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=5000)
